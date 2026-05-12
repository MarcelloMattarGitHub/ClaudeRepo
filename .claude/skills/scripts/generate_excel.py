#!/usr/bin/env python3
"""Generates formatted Excel workbook from enriched requirements JSON."""
import argparse, json, os, sys
from datetime import datetime
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl --break-system-packages"); sys.exit(1)

# Style constants
HDR_FILL = PatternFill("solid", fgColor="032D60")
HDR_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Arial", size=9)
BOLD_FONT = Font(name="Arial", bold=True, size=9)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN_BORDER = Border(
    left=Side(style="thin", color="DDDDDD"), right=Side(style="thin", color="DDDDDD"),
    top=Side(style="thin", color="DDDDDD"), bottom=Side(style="thin", color="DDDDDD"))

COMPLEXITY_FILLS = {
    "Simple": PatternFill("solid", fgColor="E8F5E9"),
    "Medium": PatternFill("solid", fgColor="FFF3E0"),
    "Complex": PatternFill("solid", fgColor="FFEBEE"),
    "Epic": PatternFill("solid", fgColor="F3E5F5"),
}
PRIORITY_FILLS = {
    "high": PatternFill("solid", fgColor="FFCDD2"),
    "medium": PatternFill("solid", fgColor="FFF9C4"),
    "low": PatternFill("solid", fgColor="E8F5E9"),
}

def style_header(ws, row, col_count):
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HDR_FILL; cell.font = HDR_FONT; cell.alignment = WRAP; cell.border = THIN_BORDER

def style_cell(cell, wrap=True):
    cell.font = BODY_FONT; cell.border = THIN_BORDER
    if wrap: cell.alignment = WRAP

def auto_width(ws, min_w=12, max_w=50):
    for col_cells in ws.columns:
        col_letter = get_column_letter(col_cells[0].column)
        max_len = max((len(str(c.value or "")) for c in col_cells), default=0)
        ws.column_dimensions[col_letter].width = min(max(max_len * 1.1, min_w), max_w)

def build_stories_sheet(wb, stories):
    ws = wb.active; ws.title = "User Stories"
    headers = ["Req ID", "Section", "Priority", "Requirement", "User Story",
               "SF Component", "Complexity", "Story Points", "Sprint",
               "Acceptance Criteria", "Dependencies", "Notes"]
    ws.append(headers); style_header(ws, 1, len(headers))
    for s in stories:
        row = [s.get("id",""), s.get("section",""), s.get("priority",""),
               s.get("text",""), s.get("user_story",""), s.get("sf_component",""),
               s.get("complexity",""), s.get("story_points",""), s.get("sprint",""),
               s.get("acceptance_criteria",""), s.get("dependencies",""), s.get("notes","")]
        ws.append(row)
        r = ws.max_row
        for c in range(1, len(headers)+1): style_cell(ws.cell(r, c))
        # Color complexity
        complexity = s.get("complexity", "")
        if complexity in COMPLEXITY_FILLS: ws.cell(r, 7).fill = COMPLEXITY_FILLS[complexity]
        # Color priority
        priority = s.get("priority", "")
        if priority in PRIORITY_FILLS: ws.cell(r, 3).fill = PRIORITY_FILLS[priority]
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    ws.freeze_panes = "A2"
    # Set specific column widths
    widths = {"A": 10, "B": 18, "C": 10, "D": 40, "E": 40, "F": 22,
              "G": 12, "H": 8, "I": 10, "J": 45, "K": 18, "L": 20}
    for col, w in widths.items(): ws.column_dimensions[col].width = w
    ws.sheet_properties.tabColor = "0176D3"

def build_summary_sheet(wb, stories):
    ws = wb.create_sheet("Effort Summary")
    sprints = {}
    for s in stories:
        sp = s.get("sprint", "Unassigned")
        if sp not in sprints: sprints[sp] = {"stories": 0, "simple": 0, "medium": 0, "complex": 0, "epic": 0, "sp": 0}
        sprints[sp]["stories"] += 1
        c = s.get("complexity", "").lower()
        if c in sprints[sp]: sprints[sp][c] += 1
        sprints[sp]["sp"] += s.get("story_points", 0) if isinstance(s.get("story_points"), (int, float)) else 0
    headers = ["Sprint", "Total Stories", "Simple", "Medium", "Complex", "Epic", "Total SP", "Est. Days (velocity=20)"]
    ws.append(headers); style_header(ws, 1, len(headers))
    for sp_name, d in sorted(sprints.items()):
        est_days = round(d["sp"] / 20 * 10, 1) if d["sp"] > 0 else 0  # 20 SP per sprint, 10 working days
        ws.append([sp_name, d["stories"], d["simple"], d["medium"], d["complex"], d["epic"], d["sp"], est_days])
        r = ws.max_row
        for c in range(1, len(headers)+1): style_cell(ws.cell(r, c))
    # Totals row
    total_row = ws.max_row + 1
    ws.cell(total_row, 1, "TOTAL").font = BOLD_FONT
    for col in range(2, len(headers)+1):
        ws.cell(total_row, col).font = BOLD_FONT
        ws.cell(total_row, col, f"=SUM({get_column_letter(col)}2:{get_column_letter(col)}{total_row-1})")
    auto_width(ws)
    ws.sheet_properties.tabColor = "2E844A"

def build_component_sheet(wb, stories):
    ws = wb.create_sheet("Component Breakdown")
    comps = {}
    for s in stories:
        c = s.get("sf_component", "Unknown")
        for part in c.split(","):
            part = part.strip()
            if not part: continue
            if part not in comps: comps[part] = {"count": 0, "sp": 0, "examples": []}
            comps[part]["count"] += 1
            sp_val = s.get("story_points", 0)
            comps[part]["sp"] += sp_val if isinstance(sp_val, (int, float)) else 0
            if len(comps[part]["examples"]) < 3: comps[part]["examples"].append(s.get("id", ""))
    headers = ["Component Type", "Count", "Total SP", "Example Reqs"]
    ws.append(headers); style_header(ws, 1, len(headers))
    for comp, d in sorted(comps.items(), key=lambda x: -x[1]["count"]):
        ws.append([comp, d["count"], d["sp"], ", ".join(d["examples"])])
        r = ws.max_row
        for c in range(1, len(headers)+1): style_cell(ws.cell(r, c))
    auto_width(ws)
    ws.sheet_properties.tabColor = "9050E9"

def main():
    ap = argparse.ArgumentParser(description="Generate Excel from enriched requirements JSON")
    ap.add_argument("--input", required=True); ap.add_argument("--output", required=True)
    args = ap.parse_args()
    if not os.path.exists(args.input): print(f"ERROR: Not found: {args.input}"); sys.exit(1)
    with open(args.input) as f: data = json.load(f)
    stories = data.get("requirements", data.get("stories", []))
    os.makedirs(os.path.dirname(args.output) or ".", exist_ok=True)
    print(f"\n{'='*60}\nExcel Generator\n{'='*60}")
    print(f"  Stories: {len(stories)}")
    wb = Workbook()
    build_stories_sheet(wb, stories)
    build_summary_sheet(wb, stories)
    build_component_sheet(wb, stories)
    wb.save(args.output)
    print(f"  Output: {args.output} ({os.path.getsize(args.output):,} bytes)")
    print(f"{'='*60}\n  Done!\n{'='*60}\n")

if __name__ == "__main__": main()
